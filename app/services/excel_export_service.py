from datetime import date, datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from app import models


class ExcelExportService:
    def build_workbook_bytes(self, db: Session, *, user_id: int, channel_id: int | None = None) -> bytes:
        workbook = Workbook()
        workbook.remove(workbook.active)

        channels = self._load_channels(db=db, user_id=user_id, channel_id=channel_id)
        channel_map = {c.id: c for c in channels}

        ideas = self._load_ideas(db=db, user_id=user_id, channel_id=channel_id)
        scripts = self._load_scripts(db=db, user_id=user_id, channel_id=channel_id)
        seo_assets = self._load_seo_assets(db=db, user_id=user_id, channel_id=channel_id)
        calendar_items = self._load_calendar_items(db=db, user_id=user_id, channel_id=channel_id)

        self._build_content_history_sheet(workbook, channel_map=channel_map, ideas=ideas, scripts=scripts, seo_assets=seo_assets)
        self._build_analysis_sheet(workbook, channels=channels)
        self._build_new_ideas_sheet(workbook, channel_map=channel_map, ideas=ideas)
        self._build_calendar_sheet(workbook, channel_map=channel_map, calendar_items=calendar_items)
        self._build_summary_sheet(
            workbook,
            channel_count=len(channels),
            ideas_count=len(ideas),
            scripts_count=len(scripts),
            seo_count=len(seo_assets),
            calendar_count=len(calendar_items),
        )

        out = BytesIO()
        workbook.save(out)
        out.seek(0)
        return out.read()

    def _build_content_history_sheet(self, wb: Workbook, *, channel_map: dict[int, models.Channel], ideas, scripts, seo_assets) -> None:
        ws = wb.create_sheet("Content History")
        headers = ["Type", "Channel", "Primary Text", "Secondary Text", "Created At"]
        self._append_header(ws, headers)

        rows = []
        for idea in ideas:
            rows.append(
                (
                    idea.created_at,
                    [
                        "Idea",
                        self._channel_name(channel_map, idea.channel_id),
                        idea.title,
                        idea.hook,
                        self._format_datetime(idea.created_at),
                    ],
                )
            )
        for script in scripts:
            rows.append(
                (
                    script.created_at,
                    [
                        "Script",
                        self._channel_name(channel_map, script.channel_id),
                        script.script_text,
                        script.cta,
                        self._format_datetime(script.created_at),
                    ],
                )
            )
        for seo in seo_assets:
            rows.append(
                (
                    seo.created_at,
                    [
                        "SEO",
                        self._channel_name(channel_map, seo.channel_id),
                        seo.title,
                        seo.description,
                        self._format_datetime(seo.created_at),
                    ],
                )
            )
        for _, row in sorted(rows, key=lambda item: item[0], reverse=True):
            ws.append(row)
        self._finalize_sheet(ws, widths={1: 16, 2: 24, 3: 52, 4: 64, 5: 18})

    def _build_analysis_sheet(self, wb: Workbook, *, channels: list[models.Channel]) -> None:
        ws = wb.create_sheet("AI Channel Analysis")
        headers = ["Channel", "Niche", "Last Analyzed", "Summary", "Strengths", "Opportunities", "Next Actions"]
        self._append_header(ws, headers)

        for channel in channels:
            payload = channel.analysis_payload or {}
            ws.append(
                [
                    channel.name,
                    channel.niche or "",
                    self._format_datetime(channel.last_analyzed_at),
                    payload.get("summary", ""),
                    ", ".join(payload.get("strengths", [])),
                    ", ".join(payload.get("opportunities", [])),
                    ", ".join(payload.get("next_actions", [])),
                ]
            )
        self._finalize_sheet(ws, widths={1: 24, 2: 22, 3: 18, 4: 52, 5: 42, 6: 42, 7: 42})

    def _build_new_ideas_sheet(self, wb: Workbook, *, channel_map: dict[int, models.Channel], ideas: list[models.ContentIdea]) -> None:
        ws = wb.create_sheet("New Content Ideas")
        headers = ["Channel", "Title", "Hook", "Angle", "Virality Score", "Status", "Created At"]
        self._append_header(ws, headers)

        for idea in ideas:
            ws.append(
                [
                    self._channel_name(channel_map, idea.channel_id),
                    idea.title,
                    idea.hook,
                    idea.angle,
                    idea.estimated_virality_score or "",
                    idea.status,
                    self._format_datetime(idea.created_at),
                ]
            )
        self._finalize_sheet(ws, widths={1: 24, 2: 42, 3: 52, 4: 52, 5: 16, 6: 14, 7: 18})

    def _build_calendar_sheet(
        self,
        wb: Workbook,
        *,
        channel_map: dict[int, models.Channel],
        calendar_items: list[models.ContentCalendar],
    ) -> None:
        ws = wb.create_sheet("Content Calendar")
        headers = ["Channel", "Publish Date", "Platform", "Status", "Notes", "Created At"]
        self._append_header(ws, headers)

        for item in calendar_items:
            ws.append(
                [
                    self._channel_name(channel_map, item.channel_id),
                    item.publish_date.isoformat() if isinstance(item.publish_date, date) else item.publish_date,
                    item.platform,
                    item.status,
                    item.notes or "",
                    self._format_datetime(item.created_at),
                ]
            )
        self._finalize_sheet(ws, widths={1: 24, 2: 16, 3: 18, 4: 14, 5: 48, 6: 18})

    def _build_summary_sheet(
        self,
        wb: Workbook,
        *,
        channel_count: int,
        ideas_count: int,
        scripts_count: int,
        seo_count: int,
        calendar_count: int,
    ) -> None:
        ws = wb.create_sheet("Performance Summary")
        headers = ["Metric", "Value"]
        self._append_header(ws, headers)
        ws.append(["Total Channels", channel_count])
        ws.append(["Total Ideas", ideas_count])
        ws.append(["Total Scripts", scripts_count])
        ws.append(["Total SEO Assets", seo_count])
        ws.append(["Calendar Entries", calendar_count])
        ws.append(["Export Version", "MVP v0.1"])
        self._finalize_sheet(ws, widths={1: 28, 2: 16})

    @staticmethod
    def _append_header(ws, headers: list[str]) -> None:
        ws.append(headers)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        fill = PatternFill("solid", fgColor="1F2937")
        font = Font(bold=True, color="FFFFFF")
        border = Border(bottom=Side(style="thin", color="D1D5DB"))
        for cell in ws[1]:
            cell.fill = fill
            cell.font = font
            cell.border = border
            cell.alignment = Alignment(vertical="center")

    @staticmethod
    def _finalize_sheet(ws, *, widths: dict[int, int] | None = None) -> None:
        ws.auto_filter.ref = ws.dimensions
        for col in ws.columns:
            values = [str(cell.value) for cell in col if cell.value is not None]
            max_len = max((len(v) for v in values), default=10)
            column_index = col[0].column
            default_width = min(max(max_len + 2, 12), 60)
            ws.column_dimensions[get_column_letter(column_index)].width = (widths or {}).get(column_index, default_width)
        for row in ws.iter_rows():
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)

    @staticmethod
    def _format_datetime(value: datetime | None) -> str:
        if not value:
            return ""
        return value.strftime("%Y-%m-%d %H:%M")

    @staticmethod
    def _channel_name(channel_map: dict[int, models.Channel], channel_id: int) -> str:
        channel = channel_map.get(channel_id)
        return channel.name if channel else str(channel_id)

    @staticmethod
    def _load_channels(db: Session, *, user_id: int, channel_id: int | None) -> list[models.Channel]:
        query = db.query(models.Channel).filter(models.Channel.user_id == user_id)
        if channel_id:
            query = query.filter(models.Channel.id == channel_id)
        return query.order_by(models.Channel.created_at.desc()).all()

    @staticmethod
    def _load_ideas(db: Session, *, user_id: int, channel_id: int | None) -> list[models.ContentIdea]:
        query = db.query(models.ContentIdea).filter(models.ContentIdea.user_id == user_id)
        if channel_id:
            query = query.filter(models.ContentIdea.channel_id == channel_id)
        return query.order_by(models.ContentIdea.created_at.desc()).all()

    @staticmethod
    def _load_scripts(db: Session, *, user_id: int, channel_id: int | None) -> list[models.Script]:
        query = db.query(models.Script).filter(models.Script.user_id == user_id)
        if channel_id:
            query = query.filter(models.Script.channel_id == channel_id)
        return query.order_by(models.Script.created_at.desc()).all()

    @staticmethod
    def _load_seo_assets(db: Session, *, user_id: int, channel_id: int | None) -> list[models.SeoAsset]:
        query = db.query(models.SeoAsset).filter(models.SeoAsset.user_id == user_id)
        if channel_id:
            query = query.filter(models.SeoAsset.channel_id == channel_id)
        return query.order_by(models.SeoAsset.created_at.desc()).all()

    @staticmethod
    def _load_calendar_items(db: Session, *, user_id: int, channel_id: int | None) -> list[models.ContentCalendar]:
        query = db.query(models.ContentCalendar).filter(models.ContentCalendar.user_id == user_id)
        if channel_id:
            query = query.filter(models.ContentCalendar.channel_id == channel_id)
        return query.order_by(models.ContentCalendar.publish_date.asc()).all()
