from io import BytesIO

from fastapi.testclient import TestClient
from openpyxl import load_workbook


def test_mvp_smoke_flow(monkeypatch, tmp_path):
    monkeypatch.setenv("DATABASE_URL", f"sqlite:///{tmp_path / 'creatorflow_test.db'}")
    monkeypatch.setenv("AUTO_CREATE_TABLES", "true")
    monkeypatch.delenv("GROQ_API_KEY", raising=False)

    from app.main import app

    with TestClient(app) as client:
        health = client.get("/health")
        assert health.status_code == 200
        assert health.json()["success"] is True

        channel_response = client.post(
            "/api/channels",
            json={
                "name": "CreatorFlow Lab",
                "platform": "youtube",
                "niche": "AI creator tools",
                "audience": "US",
                "tone": "dramatic",
                "content_type": "shorts",
                "youtube_handle": "@creatorflowlab",
                "language": "en",
            },
        )
        assert channel_response.status_code == 201
        channel_id = channel_response.json()["data"]["channel"]["id"]

        ideas_response = client.post(
            "/api/ai/generate-ideas",
            json={"channel_id": channel_id, "topic": "Shorts automation", "count": 2},
        )
        assert ideas_response.status_code == 200
        ideas_data = ideas_response.json()["data"]
        assert len(ideas_data["ideas"]) == 2

        script_response = client.post(
            "/api/ai/generate-script",
            json={
                "channel_id": channel_id,
                "idea_title": ideas_data["ideas"][0]["title"],
                "hook": ideas_data["ideas"][0]["hook"],
                "duration_seconds": 30,
                "style": "fast-paced",
                "target_audience": "solo creators",
            },
        )
        assert script_response.status_code == 200
        script_data = script_response.json()["data"]
        assert script_data["script_id"] is not None

        seo_response = client.post(
            "/api/ai/generate-seo",
            json={
                "channel_id": channel_id,
                "script_id": script_data["script_id"],
                "script_text": script_data["script_text"],
                "title": ideas_data["ideas"][0]["title"],
                "target_keywords": ["youtube shorts automation", "ai creator tools"],
                "platform": "youtube",
            },
        )
        assert seo_response.status_code == 200
        assert seo_response.json()["data"]["seo_asset_id"] is not None

        export_response = client.get("/api/export/excel")
        assert export_response.status_code == 200
        workbook = load_workbook(BytesIO(export_response.content))
        assert workbook.sheetnames == [
            "Content History",
            "AI Channel Analysis",
            "New Content Ideas",
            "Content Calendar",
            "Performance Summary",
        ]
        assert workbook["Content History"].freeze_panes == "A2"
        assert workbook["Content History"].auto_filter.ref is not None
